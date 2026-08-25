import concurrent.futures
import logging
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from strava_competition.errors import ExcelFormatError
from strava_competition.excel_writer import (
    RUNNERS_SHEET,
    STRAVA_ID_COLUMN,
    REFRESH_TOKEN_COLUMN,
    SEGMENT_TEAM_COLUMN,
    DISTANCE_TEAM_COLUMN,
    BIRTHDAY_COLUMN,
    _atomic_update_workbook,
    update_runner_refresh_tokens,
    update_single_runner_refresh_token,
)
from strava_competition.models import Runner
from typing import Any


@pytest.fixture
def runners_sheet(tmp_path: Path) -> Path:
    path = Path(tmp_path) / "tokens.xlsx"
    df = pd.DataFrame(
        {
            "Name": ["Ana", "Ben", "Cara"],
            STRAVA_ID_COLUMN: ["101", "202", "303"],
            REFRESH_TOKEN_COLUMN: ["tok1", "tok2", "tok3"],
            SEGMENT_TEAM_COLUMN: [None, None, None],
            DISTANCE_TEAM_COLUMN: [None, None, None],
            BIRTHDAY_COLUMN: [
                pd.Timestamp("2001-05-10"),
                "3-Nov",
                pd.Timestamp("1999-01-01 08:00:00"),
            ],
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=RUNNERS_SHEET, index=False)
    return path


def _refresh_token_for(df: pd.DataFrame, strava_id: str) -> Any:
    ids = df[STRAVA_ID_COLUMN].astype(str).str.strip()
    series = df.loc[ids == str(strava_id), REFRESH_TOKEN_COLUMN]
    assert not series.empty
    return series.iat[0]


def test_update_runner_refresh_tokens_updates_all_rows(runners_sheet: Any) -> None:
    runners = [
        Runner("Ana", "101", "new-ana"),
        Runner("Ben", "202", "new-ben"),
        Runner("Cara", "303", "new-cara"),
    ]

    update_runner_refresh_tokens(str(runners_sheet), runners)

    result = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    assert _refresh_token_for(result, "101") == "new-ana"
    assert _refresh_token_for(result, "202") == "new-ben"
    assert _refresh_token_for(result, "303") == "new-cara"


def test_update_runner_refresh_tokens_preserves_other_cells(runners_sheet: Any) -> None:
    """Only token cells change; birthdays and other columns are untouched."""
    before = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)

    update_runner_refresh_tokens(str(runners_sheet), [Runner("Ana", "101", "tok-new")])

    after = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    assert _refresh_token_for(after, "101") == "tok-new"
    assert _refresh_token_for(after, "202") == "tok2"
    assert after["Name"].tolist() == before["Name"].tolist()
    assert after[BIRTHDAY_COLUMN].tolist() == before[BIRTHDAY_COLUMN].tolist()


def test_update_single_runner_refresh_token_threadsafe(runners_sheet: Any) -> None:
    runners = [
        Runner("Ana", "101", "tok-ana-1"),
        Runner("Ben", "202", "tok-ben-1"),
        Runner("Cara", "303", "tok-cara-1"),
    ]

    def rotate(runner: Runner, suffix: str) -> None:
        update_single_runner_refresh_token(
            str(runners_sheet),
            Runner(runner.name, runner.strava_id, f"{runner.refresh_token}:{suffix}"),
        )

    with concurrent.futures.ThreadPoolExecutor(max_workers=len(runners)) as executor:
        futures = [executor.submit(rotate, runner, "thread") for runner in runners]
        concurrent.futures.wait(futures)

    result = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    for runner in runners:
        expected = f"{runner.refresh_token}:thread"
        actual = _refresh_token_for(result, str(runner.strava_id))
        assert actual == expected


def test_atomic_update_failure_preserves_workbook(runners_sheet: Any) -> None:
    """A crash mid-write must leave the original workbook untouched.

    Simulates a failure while mutating the temporary copy and asserts the
    original tokens survive and no stray temp file is left behind.
    """

    def boom(*_a: Any, **_kw: Any) -> Any:
        raise RuntimeError("simulated crash mid-write")

    with pytest.raises(RuntimeError, match="simulated crash"):
        _atomic_update_workbook(str(runners_sheet), boom)

    # Original workbook intact
    result = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    assert _refresh_token_for(result, "101") == "tok1"
    assert _refresh_token_for(result, "202") == "tok2"
    assert _refresh_token_for(result, "303") == "tok3"

    # No leftover temp files in the workbook's directory
    workbook = Path(runners_sheet)
    leftovers = [
        p
        for p in workbook.parent.iterdir()
        if p != workbook and p.name.startswith(f".{workbook.stem}.")
    ]
    assert not leftovers, f"Temp files leaked: {leftovers}"


def test_update_single_runner_token_write_failure_preserves_workbook(
    runners_sheet: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A write error during single-runner persistence must not corrupt data."""

    def boom(*_a: Any, **_kw: Any) -> Any:
        raise OSError("disk full")

    monkeypatch.setattr("strava_competition.excel_writer.load_workbook", boom)

    # Swallowed (OSError) and logged — should not raise.
    update_single_runner_refresh_token(
        str(runners_sheet), Runner("Ana", "101", "rotated-token")
    )

    result = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    assert _refresh_token_for(result, "101") == "tok1"

    workbook = Path(runners_sheet)
    leftovers = [
        p
        for p in workbook.parent.iterdir()
        if p != workbook and p.name.startswith(f".{workbook.stem}.")
    ]
    assert not leftovers, f"Temp files leaked: {leftovers}"


def test_update_runner_refresh_tokens_warns_for_unmatched_runner(
    runners_sheet: Any, caplog: pytest.LogCaptureFixture
) -> None:
    """An unmatched Strava ID logs a warning naming the runner, never the token."""
    with caplog.at_level(logging.WARNING, logger="strava_competition.excel_writer"):
        update_runner_refresh_tokens(
            str(runners_sheet),
            [Runner("Ghost", "999", "secret-token"), Runner("Ana", "101", "tok-new")],
        )

    assert "Ghost" in caplog.text
    assert "999" in caplog.text
    assert "secret-token" not in caplog.text
    # Matched runner still updated
    result = pd.read_excel(runners_sheet, sheet_name=RUNNERS_SHEET)
    assert _refresh_token_for(result, "101") == "tok-new"


def test_update_runner_refresh_tokens_missing_columns_raises(tmp_path: Path) -> None:
    """Missing required columns raise ExcelFormatError."""
    path = tmp_path / "broken.xlsx"
    df = pd.DataFrame({"Name": ["Ana"]})
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=RUNNERS_SHEET, index=False)

    path_str = str(path)
    runners = [Runner("Ana", "101", "tok")]
    with pytest.raises(ExcelFormatError, match="Missing columns"):
        update_runner_refresh_tokens(path_str, runners)


def test_update_preserves_formatting_and_int_ids_with_blanks(tmp_path: Path) -> None:
    """A formatted sheet with blank IDs round-trips without dtype mangling."""
    path = tmp_path / "formatted.xlsx"
    df = pd.DataFrame(
        {
            "Name": ["Ana", "Gap", "Cara"],
            STRAVA_ID_COLUMN: [101, None, 303],  # blank ID present
            REFRESH_TOKEN_COLUMN: ["tok1", None, "tok3"],
            SEGMENT_TEAM_COLUMN: ["Red", None, "Blue"],
            DISTANCE_TEAM_COLUMN: [None, None, None],
            BIRTHDAY_COLUMN: ["10-May", None, "01-Jan"],
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=RUNNERS_SHEET, index=False)
    # Apply bespoke formatting that a pandas round-trip would destroy.
    wb = load_workbook(path)
    ws = wb[RUNNERS_SHEET]
    ws.cell(row=1, column=1).font = Font(bold=True, italic=True)
    wb.save(path)
    wb.close()

    update_runner_refresh_tokens(
        str(path),
        [Runner("Ana", "101", "tok1-new"), Runner("Cara", "303", "tok3-new")],
    )

    wb = load_workbook(path)
    ws = wb[RUNNERS_SHEET]
    header_font = ws.cell(row=1, column=1).font
    assert header_font.bold is True
    assert header_font.italic is True
    # Strava IDs still stored as ints, not floats (blank rows must not
    # drift the column to 101.0 / 303.0).
    assert ws.cell(row=2, column=2).value == 101
    assert isinstance(ws.cell(row=2, column=2).value, int)
    assert ws.cell(row=4, column=2).value == 303
    assert isinstance(ws.cell(row=4, column=2).value, int)
    assert ws.cell(row=2, column=3).value == "tok1-new"
    assert ws.cell(row=4, column=3).value == "tok3-new"
    wb.close()
